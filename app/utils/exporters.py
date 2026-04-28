import io
import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Path to static folder relative to this file (app/utils/ -> project root -> static)
_STATIC_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'static')


def generate_pdf(test, test_questions):
    """Generate a PDF for the test using ReportLab."""
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Title'],
        fontSize=18,
        spaceAfter=6,
        alignment=TA_CENTER,
    )
    subtitle_style = ParagraphStyle(
        'Subtitle',
        parent=styles['Normal'],
        fontSize=11,
        spaceAfter=12,
        alignment=TA_CENTER,
        textColor=colors.HexColor('#555555'),
    )
    question_style = ParagraphStyle(
        'Question',
        parent=styles['Normal'],
        fontSize=11,
        spaceAfter=4,
        spaceBefore=8,
        leading=16,
    )
    option_style = ParagraphStyle(
        'Option',
        parent=styles['Normal'],
        fontSize=10,
        spaceAfter=2,
        leftIndent=20,
        leading=14,
    )
    instructions_style = ParagraphStyle(
        'Instructions',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#444444'),
        spaceAfter=12,
    )

    story = []

    story.append(Paragraph(test.title, title_style))
    story.append(Paragraph(
        f'Subject: {test.subject.name} &nbsp;|&nbsp; Duration: {test.duration_minutes} minutes'
        f' &nbsp;|&nbsp; Questions: {test.num_questions}',
        subtitle_style,
    ))
    story.append(Spacer(1, 0.5 * cm))
    story.append(Paragraph(
        'Instructions: Choose the best answer for each question. '
        'Circle the letter of your chosen answer.',
        instructions_style,
    ))
    story.append(Spacer(1, 0.3 * cm))

    for i, tq in enumerate(test_questions):
        q = tq.question

        q_text = f'<b>{i + 1}.</b> {q.question_text or "(No question text)"}'
        story.append(Paragraph(q_text, question_style))

        if q.has_image and q.image_path:
            img_full_path = os.path.join(_STATIC_DIR, q.image_path)
            if os.path.exists(img_full_path):
                try:
                    img = RLImage(img_full_path, width=8 * cm, height=4 * cm)
                    story.append(img)
                except Exception:
                    pass

        for letter, text in [('A', q.option_a), ('B', q.option_b),
                              ('C', q.option_c), ('D', q.option_d)]:
            if text:
                story.append(Paragraph(f'<b>{letter})</b> {text}', option_style))

        story.append(Spacer(1, 0.3 * cm))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_docx(test, test_questions):
    """Generate a Word document formatted as a Microsoft Forms-compatible quiz."""
    buffer = io.BytesIO()
    document = Document()

    title_para = document.add_heading(test.title, level=1)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f'Subject: {test.subject.name}  |  Duration: {test.duration_minutes} minutes'
        f'  |  Questions: {test.num_questions}'
    )
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    document.add_paragraph()

    instr = document.add_paragraph()
    instr_run = instr.add_run('Instructions: Choose the best answer for each question.')
    instr_run.font.bold = True
    instr_run.font.size = Pt(11)

    document.add_paragraph()

    for i, tq in enumerate(test_questions):
        q = tq.question

        q_para = document.add_paragraph()
        q_run = q_para.add_run(f'{i + 1}. {q.question_text or "(No question text)"}')
        q_run.font.bold = True
        q_run.font.size = Pt(11)

        if q.has_image and q.image_path:
            img_full_path = os.path.join(_STATIC_DIR, q.image_path)
            if os.path.exists(img_full_path):
                try:
                    document.add_picture(img_full_path, width=Inches(4))
                except Exception:
                    pass

        for letter, text in [('A', q.option_a), ('B', q.option_b),
                              ('C', q.option_c), ('D', q.option_d)]:
            if text:
                opt_para = document.add_paragraph(style='List Bullet')
                opt_run = opt_para.add_run(f'{letter}) {text}')
                opt_run.font.size = Pt(10)

        if q.correct_answer:
            ans_para = document.add_paragraph()
            ans_run = ans_para.add_run(f'[Answer: {q.correct_answer}]')
            ans_run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
            ans_run.font.size = Pt(9)
            ans_run.font.italic = True

        document.add_paragraph()

    document.save(buffer)
    buffer.seek(0)
    return buffer
